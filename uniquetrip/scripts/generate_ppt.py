from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()

# Helper to add a title + bullet slide
def add_bullets(title, bullets):
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line
        p.level = 0
    return slide

# Title slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "MakeMyTrip Clone – Project Review"
subtitle = slide.placeholders[1]
subtitle.text = "UniqueTrip | Node.js + MySQL + Vanilla JS"

# Overview
add_bullets("Overview", [
    "Full-stack travel booking app (flights, hotels, trains, buses, cabs, holidays)",
    "Frontend: HTML/CSS + Vanilla JS; Backend: Node.js/Express; DB: MySQL",
    "Auth with JWT; Passwords hashed with bcrypt; Rate limiting enabled",
])

# Architecture
add_bullets("Architecture", [
    "Frontend: static pages + script.js for API interactions and UI",
    "Backend: Express server (server/index.js), REST APIs",
    "Database: MySQL with bookings enhanced for multi-service",
    "Fallback: In-memory mode if DB not available (dev friendly)",
])

# Frontend Pages
add_bullets("Frontend Pages", [
    "index_new.html – Homepage with hero search and featured holidays",
    "flights_new, hotels_new, trains_new, buses_new, cabs_new – Featured cards",
    "booking.html – 3-step flow (details → payment → confirmation)",
    "my-bookings.html – All bookings with service-specific cards and badges",
])

# Booking Flows
add_bullets("Booking Flows", [
    "Flights: Book Now → booking.html?flightId=... for full details",
    "Holidays (homepage): Book Now → booking.html?serviceType=holiday&id=...",
    "Other services: modal or can be switched to details page for consistency",
])

# Backend APIs
add_bullets("Backend APIs", [
    "Auth: POST /api/auth/signup, POST /api/auth/login, GET /api/auth/profile",
    "Flights: GET /api/flights, GET /api/flights/:id, search endpoint",
    "Bookings: POST /api/bookings, GET /api/bookings, GET/DELETE by id",
    "Health: GET /api/health",
])

# Database Schema
add_bullets("Database Schema (key)", [
    "bookings: id, booking_reference, flight_id (nullable)",
    "service_type, service_id, service_name for non-flight services",
    "passenger_name/email/phone, num_passengers, total_price, status, booked_at",
])

# Security & Limits
add_bullets("Security & Limits", [
    "Passwords: bcrypt (10 salt rounds)",
    "Tokens: JWT (HS256), 2-hour expiry",
    "Rate limiting: 100 req/15 min (general), 5 attempts/15 min (auth)",
])

# My Bookings Page
add_bullets("My Bookings Page", [
    "Dynamic cards per service (flight/hotel/train/bus/cab/holiday)",
    "Color-coded badges and tailored content",
    "Filters and search (by reference or flight)",
])

# Test Results
add_bullets("Automated Tests", [
    "Authentication: 10/10 passing (signup, login, validation, JWT)",
    "Multi-service bookings: 10/10 passing (6 services + error cases)",
])

# Recent Enhancements
add_bullets("Recent Enhancements", [
    "Homepage Book Now → detailed booking page for holidays",
    "booking.html supports non-flight services with summary + payload",
    "Consistent header/footer branding (UniqueTrip)",
])

# Next Steps
add_bullets("Next Steps", [
    "Unify all services to detailed booking page flow",
    "Add service filters on My Bookings (All/Flight/Hotel/etc.)",
    "Harden validation (server and client) and consider HttpOnly JWT",
])

# Contacts
add_bullets("Contacts & URLs", [
    "App: http://localhost:3000/index_new.html",
    "My Bookings: http://localhost:3000/my-bookings.html",
    "API Health: http://localhost:3000/api/health",
])

# Save
prs.save("..\\MakeMyTrip_Clone_Project_Review.pptx")
print("Presentation generated: MakeMyTrip_Clone_Project_Review.pptx")
